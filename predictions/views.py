import csv
import logging

from django.contrib import messages
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from accounts.decorators import decideur_required
from core.services.journal import enregistrer_activite
from datasets.models import Observation
from forecasting.models import SessionEntrainement
from forecasting.services.pipeline import generer_previsions_futures
from forecasting.services.metrics import classer_modeles_par_performance
from forecasting.views import recuperer_artefacts_session

from .forms import NouvellePredictionForm
from .models import MessageAssistantIA, Prediction
from .rate_limit import enregistrer_question, limite_atteinte
from .services.assistant_ia import AssistantIndisponible, poser_question

logger = logging.getLogger(__name__)


@decideur_required
def nouvelle_prediction(request):
    session = SessionEntrainement.objects.filter(est_active=True).first()
    if session is None:
        messages.info(request, "Aucun modèle entraîné n'est disponible pour le moment.")
        return render(request, "predictions/nouvelle.html", {"session": None})

    chapitres_disponibles = sorted(
        set(Observation.objects.filter(dataset=session.dataset).values_list("chapitre", flat=True))
    )
    modeles_disponibles = [e.nom_modele for e in session.evaluations.all()]

    resultat_prevision = None

    if request.method == "POST":
        form = NouvellePredictionForm(
            request.POST, chapitres_disponibles=chapitres_disponibles, modeles_disponibles=modeles_disponibles
        )
        if form.is_valid():
            artefacts = recuperer_artefacts_session(session)
            if artefacts is None:
                messages.error(
                    request,
                    "Les modèles entraînés ne sont plus disponibles en mémoire (le serveur a peut-être "
                    "redémarré). Veuillez relancer l'entraînement avant de générer une nouvelle prévision.",
                )
                return redirect("forecasting:entrainement")

            chapitre = int(form.cleaned_data["chapitre"])
            horizon = form.cleaned_data["horizon"]
            nom_modele = form.cleaned_data.get("modele") or None
            mois_historique_limite = form.cleaned_data.get("mois_historique_limite")

            df_prevision = None
            modele_effectivement_utilise = None
            erreurs_par_modele = {}

            if nom_modele:
                # Modèle explicitement choisi par l'utilisateur : pas de repli
                # automatique, on respecte son choix et on affiche l'erreur
                # précise si ça échoue (elle est déjà rédigée clairement par
                # les services de modélisation, ex: "12 derniers mois non
                # tous observés, essayez tel autre modèle").
                try:
                    df_prevision = generer_previsions_futures(
                        artefacts, chapitre, horizon, nom_modele=nom_modele,
                        mois_historique_limite=mois_historique_limite,
                    )
                    modele_effectivement_utilise = nom_modele
                except Exception as exc:
                    logger.exception("Échec de génération de la prévision (modèle explicite)")
                    messages.error(request, f"Impossible de générer la prévision avec {nom_modele} : {exc}")
                    return render(request, "predictions/nouvelle.html", {"form": form, "session": session})
            else:
                # Aucun modèle imposé : on essaie le meilleur, puis on se
                # replie automatiquement sur le suivant du classement si le
                # meilleur échoue pour CE chapitre précis (ex: historique
                # récent incomplet, spécifique au LSTM) — plutôt que d'échouer
                # sèchement alors qu'un autre modèle aurait très bien pu fonctionner.
                classement = classer_modeles_par_performance(artefacts.metriques)
                for candidat in classement:
                    try:
                        df_prevision = generer_previsions_futures(
                            artefacts, chapitre, horizon, nom_modele=candidat,
                            mois_historique_limite=mois_historique_limite,
                        )
                        modele_effectivement_utilise = candidat
                        break
                    except Exception as exc:
                        erreurs_par_modele[candidat] = str(exc)
                        continue

                if df_prevision is None:
                    logger.error("Échec de génération de la prévision pour tous les modèles : %s", erreurs_par_modele)
                    messages.error(
                        request,
                        "Aucun des modèles disponibles n'a pu générer de prévision pour ce chapitre. "
                        "L'historique disponible est probablement trop incomplet pour cet horizon.",
                    )
                    return render(request, "predictions/nouvelle.html", {"form": form, "session": session})

                if modele_effectivement_utilise != classement[0]:
                    messages.info(
                        request,
                        f"Le meilleur modèle ({classement[0]}) n'a pas pu être utilisé pour ce chapitre "
                        f"({erreurs_par_modele.get(classement[0], 'historique insuffisant')}) — "
                        f"{modele_effectivement_utilise} a été utilisé à la place.",
                    )

            for avertissement in df_prevision.attrs.get("avertissements", []):
                messages.warning(request, avertissement)

            modele_final = df_prevision["modele"].iloc[0]
            evaluation = session.evaluations.filter(nom_modele=modele_final).first()

            prediction = Prediction.objects.create(
                utilisateur=request.user,
                session_entrainement=session,
                chapitre=chapitre,
                modele_utilise=modele_final,
                horizon_mois=len(df_prevision),
                dates_predites=[d.strftime("%m-%Y") for d in df_prevision["date"]],
                valeurs_predites=[round(float(v), 6) for v in df_prevision["taux_prevu"]],
                mae_modele=evaluation.mae if evaluation else None,
                rmse_modele=evaluation.rmse if evaluation else None,
                r2_modele=evaluation.r2 if evaluation else None,
                mape_modele=evaluation.mape if evaluation else None,
            )
            enregistrer_activite(
                request.user,
                f"Prévision générée — chapitre {chapitre}, {horizon} mois, modèle {modele_final}",
                request=request,
            )
            messages.success(request, "Prévision générée avec succès.")
            resultat_prevision = prediction
        else:
            messages.error(request, "Le formulaire contient des erreurs. Veuillez les corriger.")
    else:
        form = NouvellePredictionForm(
            chapitres_disponibles=chapitres_disponibles, modeles_disponibles=modeles_disponibles
        )

    # Historique récent des observations du chapitre sélectionné (pour le graphique réel + prévu)
    historique_recent = []
    messages_ia = []
    if resultat_prevision:
        obs = (
            Observation.objects.filter(dataset=session.dataset, chapitre=resultat_prevision.chapitre)
            .order_by("date_mois")
        )
        historique_recent = [
            {"date": o.date_mois.strftime("%m-%Y"), "taux": o.taux_engagement_ae} for o in obs
        ]
        messages_ia = resultat_prevision.messages_ia.all()

    return render(
        request,
        "predictions/nouvelle.html",
        {
            "form": form,
            "session": session,
            "resultat_prevision": resultat_prevision,
            "historique_recent": historique_recent,
            "messages_ia": messages_ia,
        },
    )


@decideur_required
def historique_predictions(request):
    predictions = Prediction.objects.select_related("utilisateur", "session_entrainement").all()

    chapitre = request.GET.get("chapitre", "").strip()
    modele = request.GET.get("modele", "").strip()
    if chapitre.isdigit():
        predictions = predictions.filter(chapitre=int(chapitre))
    if modele:
        predictions = predictions.filter(modele_utilise=modele)

    if request.GET.get("export") == "csv":
        return _exporter_predictions_csv(predictions)
    if request.GET.get("export") == "xlsx":
        return _exporter_predictions_xlsx(predictions)

    paginator = Paginator(predictions, 15)
    page = paginator.get_page(request.GET.get("page"))

    return render(
        request,
        "predictions/historique.html",
        {"page_obj": page, "chapitre_filtre": chapitre, "modele_filtre": modele},
    )


@decideur_required
def exporter_prediction(request, pk, format_export):
    """Export d'UNE prévision individuelle, au format CSV ou Excel."""
    prediction = get_object_or_404(Prediction, pk=pk)
    queryset = Prediction.objects.filter(pk=prediction.pk)

    if format_export == "xlsx":
        return _exporter_predictions_xlsx(
            queryset, nom_fichier=f"prevision_chapitre_{prediction.chapitre}.xlsx"
        )
    if format_export == "csv":
        response = _exporter_predictions_csv(queryset)
        response["Content-Disposition"] = (
            f'attachment; filename="prevision_chapitre_{prediction.chapitre}.csv"'
        )
        return response

    messages.error(request, "Format d'export non reconnu.")
    return redirect("predictions:historique")


@decideur_required
@require_POST
def poser_question_ia(request, pk):
    """Endpoint AJAX : reçoit une question du décideur sur une prévision précise,
    interroge l'assistant IA avec le contexte exact de cette prévision, et
    retourne la réponse en JSON pour affichage dans le widget de chat.
    """
    prediction = get_object_or_404(Prediction, pk=pk)

    if limite_atteinte(request.user):
        return JsonResponse(
            {"erreur": "Vous avez atteint la limite de questions à l'assistant IA pour cette heure. Réessayez plus tard."},
            status=429,
        )

    question = request.POST.get("question", "").strip()
    if not question:
        return JsonResponse({"erreur": "Veuillez saisir une question."}, status=400)
    if len(question) > 500:
        return JsonResponse({"erreur": "Votre question est trop longue (500 caractères maximum)."}, status=400)

    dataset = prediction.session_entrainement.dataset
    historique = [
        {"date": o.date_mois.strftime("%m-%Y"), "taux": o.taux_engagement_ae}
        for o in Observation.objects.filter(dataset=dataset, chapitre=prediction.chapitre).order_by("date_mois")
    ]

    enregistrer_question(request.user)

    try:
        reponse_texte = poser_question(prediction, question, historique)
    except AssistantIndisponible as exc:
        return JsonResponse({"erreur": str(exc)}, status=503)

    MessageAssistantIA.objects.create(
        prediction=prediction, utilisateur=request.user, question=question, reponse=reponse_texte
    )
    enregistrer_activite(request.user, f"Question à l'assistant IA sur la prévision #{prediction.pk}", request=request)

    return JsonResponse({"reponse": reponse_texte})


@decideur_required
def supprimer_prediction(request, pk):
    prediction = get_object_or_404(Prediction, pk=pk)
    # Seul l'auteur ou un administrateur peut supprimer une prévision.
    if prediction.utilisateur != request.user and request.user.role != "ADMIN":
        messages.error(request, "Vous n'êtes pas autorisé à supprimer cette prévision.")
        return redirect("predictions:historique")

    if request.method == "POST":
        prediction.delete()
        messages.success(request, "Prévision supprimée avec succès.")
        enregistrer_activite(request.user, f"Suppression de la prévision #{pk}", request=request)
        return redirect("predictions:historique")

    return render(
        request,
        "accounts/confirmer_action.html",
        {
            "titre": "Supprimer la prévision",
            "message": "Cette action est irréversible. Confirmez-vous la suppression ?",
            "url_retour": "predictions:historique",
            "danger": True,
        },
    )


def _exporter_predictions_csv(queryset):
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="historique_previsions.csv"'
    writer = csv.writer(response)
    writer.writerow(["Date", "Chapitre", "Modele", "Taux_prevu"])
    for pred in queryset.iterator():
        for date_str, valeur in pred.lignes():
            writer.writerow([date_str, pred.chapitre, pred.modele_utilise, round(valeur, 6)])
    return response


def _exporter_predictions_xlsx(queryset, nom_fichier: str = "historique_previsions.xlsx"):
    """Génère un vrai classeur Excel (.xlsx) avec en-têtes mis en forme,
    largeurs de colonnes ajustées et pourcentages correctement formatés —
    pas un simple CSV renommé.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Prévisions AE"

    entetes = ["Date", "Chapitre", "Modèle", "Taux prévu (%)"]
    feuille.append(entetes)

    remplissage_entete = PatternFill(start_color="B7D64A", end_color="B7D64A", fill_type="solid")
    police_entete = Font(bold=True, color="22281B")
    for col_idx, _ in enumerate(entetes, start=1):
        cellule = feuille.cell(row=1, column=col_idx)
        cellule.fill = remplissage_entete
        cellule.font = police_entete
        cellule.alignment = Alignment(horizontal="center")

    ligne = 2
    for pred in queryset.iterator():
        for date_str, valeur in pred.lignes():
            feuille.cell(row=ligne, column=1, value=date_str)
            feuille.cell(row=ligne, column=2, value=f"Chapitre {pred.chapitre}")
            feuille.cell(row=ligne, column=3, value=pred.modele_utilise)
            cellule_taux = feuille.cell(row=ligne, column=4, value=round(valeur * 100, 4))
            cellule_taux.number_format = "0.00\\%"
            ligne += 1

    for col_idx, largeur in enumerate([14, 14, 14, 16], start=1):
        feuille.column_dimensions[get_column_letter(col_idx)].width = largeur

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'
    classeur.save(response)
    return response
